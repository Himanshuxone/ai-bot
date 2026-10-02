"""
Safe document loaders for financial files.

Supported formats: PDF, XLSX, CSV, TXT / Markdown.

Every loader treats the file as hostile input:

* the declared extension must match the file's magic bytes / structure;
* sizes, page counts and row counts are capped;
* XLSX archives are inspected for decompression bombs and macros before
  being parsed, and formulas are never evaluated (cached values only);
* encrypted PDFs are rejected rather than brute-forced or partially read.

[RULE: SEC-FILE-UPLOAD]       Magic bytes, size caps, zip-bomb + macro checks.
[RULE: SEC-INPUT-VALIDATION]  Strict UTF-8 decoding, control-char stripping.
[RULE: OWASP-LLM04]           Malformed / malicious files rejected at the door.
[RULE: OWASP-LLM10]           Bounded parsing cost per upload.
[RULE: EUAIA-ART10]           Provenance (file, page, sheet, row) kept with
                              every section for traceable answers.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath

from finrag.config import Settings
from finrag.security.guardrails import normalise_text


class UnsupportedDocumentError(ValueError):
    """Raised when a file fails validation and must not be ingested."""


@dataclass
class Section:
    """A logical unit of a document (a PDF page, a sheet, a text file)."""

    location: str  # human-readable provenance, e.g. "page 3" or "sheet 'P&L'"
    kind: str  # "text" or "table"
    text: str = ""  # used when kind == "text"
    header: str = ""  # used when kind == "table"
    rows: list[str] = field(default_factory=list)  # used when kind == "table"


@dataclass
class LoadedDocument:
    """Result of loading: sanitised filename, detected format and sections."""

    filename: str
    file_type: str
    sections: list[Section]


_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._ -]")


def sanitize_filename(name: str) -> str:
    """Strip directories and unsafe characters from an uploaded filename.

    [RULE: SEC-FILE-UPLOAD] Prevents path traversal ("../../etc/passwd") and
    control characters in names that end up in logs and citations.
    """
    base = PurePath(name.replace("\\", "/")).name
    cleaned = _SAFE_NAME_RE.sub("_", base).strip(" .")
    return (cleaned or "document")[:128]


def _format_cell(value: object) -> str:
    """Render a spreadsheet cell value as text, keeping numbers unambiguous."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        return f"{value:,.2f}" if abs(value) >= 1 or value == 0 else f"{value:.4g}"
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    # [RULE: SEC-INPUT-VALIDATION] Strip control / invisible characters.
    return normalise_text(str(value)).strip()


def _decode_text(data: bytes) -> str:
    """Strictly decode UTF-8 text and reject binary content."""
    if b"\x00" in data:
        raise UnsupportedDocumentError("binary content in a text file")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UnsupportedDocumentError("text files must be UTF-8 encoded") from exc


# --------------------------------------------------------------------- loaders
def _load_pdf(data: bytes, settings: Settings) -> list[Section]:
    """Extract text page by page from a PDF."""
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    if not data.startswith(b"%PDF-"):
        raise UnsupportedDocumentError("file is not a valid PDF (bad magic bytes)")
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except (PdfReadError, ValueError) as exc:
        raise UnsupportedDocumentError("unreadable PDF") from exc
    if reader.is_encrypted:
        # [RULE: SEC-FILE-UPLOAD] Do not attempt to decrypt protected files.
        raise UnsupportedDocumentError("encrypted PDFs are not accepted")
    if len(reader.pages) > settings.max_pdf_pages:
        raise UnsupportedDocumentError(f"PDF exceeds {settings.max_pdf_pages} pages")
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        text = normalise_text(page.extract_text() or "").strip()
        if text:
            sections.append(Section(location=f"page {number}", kind="text", text=text))
    return sections


def _inspect_xlsx_archive(data: bytes, settings: Settings) -> None:
    """Reject zip bombs, macro-enabled workbooks and non-workbook zips."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise UnsupportedDocumentError("file is not a valid XLSX (bad zip)") from exc
    with archive:
        names = set(archive.namelist())
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise UnsupportedDocumentError("zip archive is not an Excel workbook")
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            # [RULE: SEC-FILE-UPLOAD] Macro-enabled content is refused.
            raise UnsupportedDocumentError("macro-enabled workbooks are not accepted")
        if any(n.startswith("xl/externalLinks/") for n in names):
            # External links can trigger outbound fetches in some tooling.
            raise UnsupportedDocumentError("workbooks with external links are not accepted")
        total = sum(info.file_size for info in archive.infolist())
        compressed = max(1, sum(info.compress_size for info in archive.infolist()))
        # [RULE: SEC-FILE-UPLOAD] [RULE: OWASP-LLM10] Decompression-bomb guard.
        if total > settings.max_xlsx_uncompressed_bytes:
            raise UnsupportedDocumentError("workbook expands beyond the allowed size")
        if total / compressed > settings.max_xlsx_compression_ratio:
            raise UnsupportedDocumentError("suspicious compression ratio (possible zip bomb)")


def _load_xlsx(data: bytes, settings: Settings) -> list[Section]:
    """Load every worksheet as a table section (header row + data rows)."""
    from openpyxl import load_workbook

    _inspect_xlsx_archive(data, settings)
    # read_only streams rows; data_only returns cached values and NEVER
    # evaluates formulas, so formula payloads cannot execute.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[Section] = []
    total_rows = 0
    try:
        for sheet in workbook.worksheets:
            header: list[str] | None = None
            rows: list[str] = []
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [_format_cell(v) for v in row]
                if not any(cells):
                    continue
                total_rows += 1
                if total_rows > settings.max_sheet_rows:
                    raise UnsupportedDocumentError(
                        f"workbook exceeds {settings.max_sheet_rows} rows")
                if header is None:
                    header = [c or f"col{i + 1}" for i, c in enumerate(cells)]
                    continue
                # "Row 7 | Line item: Revenue | FY2023: 1,200.00 | FY2024: 1,450.00"
                pairs = [f"{header[i] if i < len(header) else f'col{i + 1}'}: {c}"
                         for i, c in enumerate(cells) if c]
                rows.append(f"Row {row_idx} | " + " | ".join(pairs))
            if header is not None:
                sections.append(Section(location=f"sheet '{_format_cell(sheet.title)}'",
                                        kind="table", header=" | ".join(header), rows=rows))
    finally:
        workbook.close()
    return sections


def _load_csv(data: bytes, settings: Settings) -> list[Section]:
    """Load a CSV file as a single table section."""
    text = _decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header: list[str] | None = None
    rows: list[str] = []
    for row_idx, row in enumerate(reader, start=1):
        cells = [_format_cell(c) for c in row]
        if not any(cells):
            continue
        if row_idx > settings.max_sheet_rows:
            raise UnsupportedDocumentError(f"CSV exceeds {settings.max_sheet_rows} rows")
        if header is None:
            header = [c or f"col{i + 1}" for i, c in enumerate(cells)]
            continue
        pairs = [f"{header[i] if i < len(header) else f'col{i + 1}'}: {c}"
                 for i, c in enumerate(cells) if c]
        rows.append(f"Row {row_idx} | " + " | ".join(pairs))
    if header is None:
        return []
    return [Section(location="table", kind="table", header=" | ".join(header), rows=rows)]


def _load_text(data: bytes, settings: Settings) -> list[Section]:
    """Load plain text or Markdown."""
    text = normalise_text(_decode_text(data)).strip()
    return [Section(location="text", kind="text", text=text)] if text else []


_LOADERS = {
    ".pdf": ("pdf", _load_pdf),
    ".xlsx": ("xlsx", _load_xlsx),
    ".csv": ("csv", _load_csv),
    ".txt": ("text", _load_text),
    ".md": ("text", _load_text),
}


def load_document(filename: str, data: bytes, settings: Settings) -> LoadedDocument:
    """Validate and parse an uploaded file into sections.

    Raises ``UnsupportedDocumentError`` on any validation failure; callers log
    the rejection and never ingest partial content.
    """
    safe_name = sanitize_filename(filename)
    # [RULE: SEC-FILE-UPLOAD] [RULE: OWASP-LLM10] Size limit before any parsing.
    if len(data) == 0:
        raise UnsupportedDocumentError("empty file")
    if len(data) > settings.max_upload_bytes:
        raise UnsupportedDocumentError(f"file exceeds {settings.max_upload_bytes} bytes")
    suffix = PurePath(safe_name).suffix.lower()
    if suffix not in _LOADERS:
        # [RULE: SEC-FILE-UPLOAD] Allow-list of extensions (deny by default).
        raise UnsupportedDocumentError(f"unsupported file type '{suffix or 'none'}'")
    file_type, loader = _LOADERS[suffix]
    sections = loader(data, settings)
    if not sections:
        raise UnsupportedDocumentError("no extractable content (scanned PDF? try OCR first)")
    return LoadedDocument(filename=safe_name, file_type=file_type, sections=sections)
